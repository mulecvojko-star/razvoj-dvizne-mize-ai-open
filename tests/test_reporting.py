"""Real document rendering, formulas, layout and failure gates; never a model request."""

from copy import deepcopy
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from src.reporting.package import FILES, build_package
from src.reporting.source import PackageError
from src.models import Handoff
from src.orchestrator.cycles import CycleRunner

ROOT=Path(__file__).resolve().parents[1]


class ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from pypdf import PdfReader
        import openpyxl
        import fitz
        cls.reader=PdfReader;cls.openpyxl=openpyxl;cls.fitz=fitz
        cls.document=json.loads((ROOT/'examples/V01-R01-engineering-result.json').read_text())
        cls.directory=tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.directory.cleanup)
        cls.archive=build_package(cls.document,cls.directory.name)
        with ZipFile(cls.archive) as archive:
            cls.files={name:archive.read(name) for name in archive.namelist()}

    def test_zip_has_exact_requested_files_and_unchanged_source_numbers(self):
        self.assertEqual(self.archive.name,'V01-R01_Razvojni_paket.zip')
        self.assertEqual(set(self.files),set(FILES))
        with ZipFile(self.archive) as archive: self.assertIsNone(archive.testzip())
        source=json.loads(self.files['05_Izvorni_podatki.json'])
        self.assertEqual(source['handoff'],self.document)
        self.assertFalse(source['engineering_verified'])
        self.assertEqual(source['technical_status'],'changes_required')

    def test_pdfs_are_valid_a4_embedded_unicode_and_keep_technical_status(self):
        for name,data in self.files.items():
            if not name.endswith('.pdf'): continue
            with self.subTest(name=name):
                pdf=self.reader(io.BytesIO(data),strict=True)
                self.assertGreater(len(pdf.pages),0)
                text='\n'.join(page.extract_text() for page in pdf.pages)
                for ch in 'čšž': self.assertIn(ch,text)
                for label in ('V01-R01','changes_required','engineering_verified=false','Sintetični podatki'):
                    self.assertIn(label,text)
                for page in pdf.pages:
                    self.assertAlmostEqual(float(page.mediabox.width),595.276,places=1)
                    self.assertAlmostEqual(float(page.mediabox.height),841.89,places=1)
                    fonts=page['/Resources']['/Font'].get_object()
                    self.assertTrue(any('/FontFile2' in font.get_object().get('/FontDescriptor',{})
                                        for font in fonts.values()))

    def test_all_pdf_text_stays_on_page_without_overlapping_spans(self):
        for name,data in self.files.items():
            if not name.endswith('.pdf'): continue
            pdf=self.fitz.open(stream=data,filetype='pdf')
            for page in pdf:
                spans=[s for b in page.get_text('dict')['blocks'] if 'lines' in b
                       for line in b['lines'] for s in line['spans'] if s['text'].strip()]
                for index,span in enumerate(spans):
                    rect=self.fitz.Rect(span['bbox'])
                    self.assertTrue(page.rect.contains(rect),(name,page.number,span['text']))
                    for other in spans[index+1:]:
                        overlap=rect & self.fitz.Rect(other['bbox'])
                        self.assertLessEqual(overlap.get_area(),3,(name,span['text'],other['text']))

    def test_constructor_and_hydraulic_tables_graphs_and_missing_values(self):
        ctext=''.join(p.extract_text() for p in self.reader(io.BytesIO(self.files['01_Konstruktor.pdf'])).pages)
        htext=''.join(p.extract_text() for p in self.reader(io.BytesIO(self.files['02_Hidravlik.pdf'])).pages)
        self.assertIn('143,08',ctext)
        self.assertIn('64.136,45',ctext)
        self.assertIn('1.000,00',ctext)
        self.assertIn('Koordinate vpetij',ctext)
        self.assertIn('dl/dH',ctext)
        self.assertIn('manjkajo vhodni podatki',htext)
        self.assertIn('niso določeni',htext)
        for value in ('71,44','63,90','57,15','50,92','160,00','200,00','250,00','315,00'):
            self.assertIn(value,htext)
        pdf=self.fitz.open(stream=self.files['01_Konstruktor.pdf'],filetype='pdf')
        self.assertTrue(any(len(d['items'])>=10 and d['width']==2 for page in pdf for d in page.get_drawings()))

    def test_excel_formulas_unknowns_synthetic_subtotal_and_a4(self):
        wb=self.openpyxl.load_workbook(io.BytesIO(self.files['03_Ekonomist.xlsx']),data_only=False)
        ws=wb['Stroški']
        self.assertEqual(ws['G4'].value,'=IF(AND(ISNUMBER(C4),ISNUMBER(E4)),ROUND(C4*E4,2),"")')
        self.assertAlmostEqual(ws['C4'].value*ws['E4'].value,220)
        self.assertEqual(ws['G11'].value,'=IF(COUNTIF(H4:H9,"SINTETIČNO")=0,"",SUMIF(H4:H9,"SINTETIČNO",G4:G9))')
        self.assertEqual(sum(ws.cell(row,6).value or 0 for row in range(4,10)),1120)
        self.assertEqual(ws['G14'].value,'NI DOLOČENA')
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertEqual(str(ws.page_setup.paperSize),str(ws.PAPERSIZE_A4))
        self.assertEqual(ws.page_setup.orientation,'landscape')
        self.assertEqual(ws.page_setup.fitToWidth,1)
        unknown=[]
        for row in range(4,10):
            if ws.cell(row,8).value=='NEZNANO':
                unknown.append(row)
                self.assertIsNone(ws.cell(row,6).value)
                self.assertIsNone(ws.cell(row,3).value)
                self.assertIsNone(ws.cell(row,5).value)
                self.assertEqual(ws.cell(row,7).value,f'=IF(ISNUMBER(F{row}),ROUND(F{row},2),"")')
        self.assertEqual(unknown,[7,9])
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.data_type=='f':
                        self.assertTrue(self.openpyxl.formula.Tokenizer(cell.value).items)
        self.assertIn('engineering_verified=false',ws['A2'].value)
        self.assertIn('Sintetično',ws['I4'].value)

    def test_incomplete_stopped_unapproved_or_changed_units_never_produce_a_zip(self):
        stopped={'status':'stopped','handoff':deepcopy(self.document)}
        cases=[stopped]
        missing=deepcopy(self.document);missing['results'].pop('ekonomist');cases.append(missing)
        approved=deepcopy(self.document);approved['results']['recenzent']['engineering_verified']=True;cases.append(approved)
        wrong_units=deepcopy(self.document);wrong_units['results']['konstruktor']['engineering']['units']['length']='m';cases.append(wrong_units)
        wrong_force=deepcopy(self.document);wrong_force['results']['hidravlik']['engineering']['force_curve'][0]['force_N']=0;cases.append(wrong_force)
        for document in cases:
            with tempfile.TemporaryDirectory() as directory,self.subTest(document=list(document)):
                with self.assertRaises(PackageError):
                    build_package(document,directory)
                self.assertFalse(list(Path(directory).glob('*.zip')))

    def test_exported_json_can_be_rendered_again_without_losing_original_run(self):
        from src.reporting.source import load
        with tempfile.TemporaryDirectory() as directory:
            first=build_package(self.document,directory,metadata={"manifest":{"run_id":"40794650126","commit":"source-commit"}})
            with ZipFile(first) as archive:
                source=Path(directory)/'05_Izvorni_podatki.json'
                source.write_bytes(archive.read('05_Izvorni_podatki.json'))
            document,metadata=load(source)
            second=build_package(document,Path(directory)/'roundtrip',metadata=metadata)
            with ZipFile(second) as archive:
                exported=json.loads(archive.read('05_Izvorni_podatki.json'))
            self.assertEqual(exported['handoff'],self.document)
            self.assertEqual(exported['provenance']['manifest']['run_id'],'40794650126')
            self.assertEqual(exported['execution']['status'],'awaiting_approval')

    def test_renderer_failure_does_not_publish_partial_zip(self):
        with tempfile.TemporaryDirectory() as directory,patch('src.reporting.documents.make_excel',side_effect=RuntimeError('failure')):
            with self.assertRaises(PackageError): build_package(self.document,directory)
            self.assertEqual(list(Path(directory).iterdir()),[])

    def test_secrets_and_html_are_not_exported_or_executed(self):
        document=deepcopy(self.document)
        document['results']['konstruktor']['summary']='čšž <b>besedilo</b> sk-proj-SUPERSECRET1234567890 password=NEVEREXPORT'
        metadata={'OPENAI_API_KEY':'KEY-MUST-NOT-APPEAR','password':'PASSWORD-MUST-NOT-APPEAR'}
        with tempfile.TemporaryDirectory() as directory:
            archive=build_package(document,directory,metadata=metadata)
            with ZipFile(archive) as zipfile:
                source=zipfile.read('05_Izvorni_podatki.json').decode()
                pdftext=''.join(p.extract_text() for p in self.reader(io.BytesIO(zipfile.read('01_Konstruktor.pdf'))).pages)
                for secret in ('SUPERSECRET1234567890','NEVEREXPORT','KEY-MUST-NOT-APPEAR','PASSWORD-MUST-NOT-APPEAR'):
                    self.assertNotIn(secret,source+pdftext)
                self.assertIn('<b>besedilo</b>',pdftext)
        from src.reporting.documents import literal
        wb=self.openpyxl.Workbook();literal(wb.active['A1'],'=HYPERLINK("https://example.com")')
        self.assertEqual(wb.active['A1'].data_type,'s')

    def test_local_command_never_runs_agents(self):
        with tempfile.TemporaryDirectory() as directory,patch('src.orchestrator.pipeline.Orchestrator.run',side_effect=AssertionError('No agents')):
            from src.reporting.__main__ import main
            with patch.object(sys,'argv',['reporting',str(ROOT/'examples/V01-R01-engineering-result.json'),'--output-dir',directory]): main()
            self.assertTrue((Path(directory)/'V01-R01_Razvojni_paket.zip').exists())

    def test_cycle_generates_package_automatically_and_retains_completed_data_on_document_failure(self):
        inputs=deepcopy(self.document)
        inputs['results']={};inputs['open_questions']=[]
        with tempfile.TemporaryDirectory() as directory:
            runner=CycleRunner(Path(directory))
            result=runner.run(Handoff(**inputs),approved_label='V01-R01')
            record=json.loads((Path(directory)/'V01-R01.json').read_text())
            self.assertEqual(record['development_package']['status'],'created')
            self.assertTrue((Path(directory)/record['development_package']['path']).exists())
            self.assertEqual(record['handoff'],result.to_dict())
        with tempfile.TemporaryDirectory() as directory,patch('src.reporting.package.build_package',side_effect=RuntimeError('broken docs')):
            with self.assertRaises(PackageError): CycleRunner(Path(directory)).run(Handoff(**inputs),approved_label='V01-R01')
            record=json.loads((Path(directory)/'V01-R01.json').read_text())
            self.assertEqual(record['status'],'awaiting_approval')
            self.assertEqual(len(record['handoff']['results']),4)
            self.assertEqual(record['development_package']['status'],'failed')
            self.assertFalse(list(Path(directory).rglob('*.zip')))
