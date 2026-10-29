"""Offline final export from validated stored trial JSON. Reuse Task 05 A4 PDFs."""
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.workbook.properties import CalcProperties
from src.models import Handoff
from src.models.contracts import ORDER, digest
from src.engineering.contracts import compare
from src.reporting.documents import make_pdfs, pdf, table, section, paragraph, fmt, WIDTH, fonts, literal
from src.reporting.source import redact
from src.orchestrator.revision_policy import fixed_requirements
from .engineering import geometry, hydraulic, economics, findings
from .agents import cited_sources
from .presentation import document_view

FILES=('00_Povzetek_razvoja.pdf','01_Konstruktor.pdf','02_Hidravlik.pdf','03_Ekonomist.xlsx','04_Recenzent.pdf',
       '05_Primerjava_revizij.pdf','06_Komunikacija_agentov.pdf','07_Viri_in_predpostavke.pdf','08_Izvorni_podatki.json','09_API_poraba.json')


def validate(state):
    if state.get('status')!='completed' or state.get('engineering_verified') is not False:raise ValueError('Končni paket zahteva zaključen, tehnično nepotrjen cikel.')
    revisions=state['revisions'];kb=state['knowledge']
    if not 1<=len(revisions)<=min(2,kb['max_revisions']) or not kb['compatible']:raise ValueError('Neveljavno število revizij ali baza.')
    usage=state['api_usage']
    if usage['call_count']>16 or any(v['calls']>4 for v in usage['per_agent'].values()):raise ValueError('Preseženi API-klici.')
    results=[]
    for index,record in enumerate(revisions):
        h=Handoff(**record['handoff'])
        if record['status']!='completed' or h.label!=f'V01-R{index+1:02d}' or tuple(h.results)!=ORDER or fixed_requirements(h.inputs)!=state['fixed_requirements']:raise ValueError('Neveljavna identiteta, fiksne zahteve ali predaja.')
        for result in h.results.values():
            if 'effective_source_ids' in result:
                citation_data={'source_ids':result['declared_source_ids'],'requests':result['requests'],
                               'request_updates':result['request_updates'],'proposal':result['proposal']}
                locations=cited_sources(citation_data)
                if result['effective_source_ids']!=sorted(locations) or result.get('citation_locations')!=locations or set(locations)-set(kb['records']):
                    raise ValueError('Neskladna sled izrecno navedenih virov.')
        c,hv,e,r=[h.results[n]['engineering'] for n in ORDER]
        if not compare(c,geometry(h.inputs)) or not compare(hv,hydraulic(h.inputs,c,kb,hv['catalogue']['choice'])) or not compare(e,economics(h.inputs,c,hv,kb)) or not compare(r,findings(h.inputs,c,hv,e,kb)):raise ValueError('Poročilo ni skladno z determinističnimi izračuni.')
        if h.results['recenzent']['decision']!='changes_required' or any(h.results[n]['engineering_verified'] is not False for n in ORDER):raise ValueError('Lažna tehnična odobritev.')
        results.append(h)
    # Closure proof cannot be supplied by the model or forged in stored JSON.
    for task in state['tasks']:
        if task['status']=='resolved':
            if task['id']!='KB-CAPACITY' or len(results)!=2:raise ValueError('Zaprtje nima podprtega dokaza.')
            a,b=[h.results['hidravlik']['engineering']['catalogue'] for h in results]
            if (a['nominal_force_pass'] and a['pressure_pass']) or not(b['nominal_force_pass'] and b['pressure_pass']):raise ValueError('Zaprtje ni dokazano.')
    return results


def workbook(path,state,results):
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(fullCalcOnLoad=True)
    for h in results:
        ws=wb.create_sheet(h.revision+' Kosovnica');ws.append([h.label+' – delni sintetični seštevek; engineering_verified=false'])
        ws.merge_cells('A1:H1');ws.append(['COST ID','Postavka','Količina','Cena enote [EUR]','Znesek [EUR]','Status','Vir','Omejitev'])
        for row,item in enumerate(h.results['ekonomist']['engineering']['bom'],3):
            for col,value in ((1,item['id']),(2,item['description']),(6,'NEZNANO' if item['amount_EUR']is None else 'SINTETIČNO'),(7,item['source']),(8,item['quantity_basis'])):literal(ws.cell(row,col),value)
            ws.cell(row,3,item['quantity']);ws.cell(row,4,item['unit_price_EUR'])
            ws.cell(row,5,f'=IF(AND(ISNUMBER(C{row}),ISNUMBER(D{row})),ROUND(C{row}*D{row},2),"")')
        end=ws.max_row
        ws.append(['Delni sintetični seštevek',None,None,None,f'=SUMIF(F3:F{end},"SINTETIČNO",E3:E{end})'])
        ws.append(['Preverjene cene','NI PREDLOŽENIH PONUDB'])
        ws.append(['Neznane postavke',f'=COUNTIF(F3:F{end},"NEZNANO")'])
        ws.append(['Potrjena investicijska vrednost','NI DOLOČENA'])
        ws.freeze_panes='C3';ws.auto_filter.ref=f'A2:H{end}'
    ws=wb.create_sheet('Primerjava');ws.append(['Revizija','Delni sintetični seštevek [EUR]','Skupna cena naprave','Cilinder','Tlak [bar]'])
    for row,h in enumerate(results,2):
        end=2+len(h.results['ekonomist']['engineering']['bom']);choice=h.results['hidravlik']['engineering']['catalogue']['choice']
        ws.append([h.revision,f"=SUM('{h.revision} Kosovnica'!E3:E{end})",'NI DOLOČENA',choice['cylinder_id'],choice['pressure_bar']])
    alt=wb.create_sheet('Kataloške alternative');alt.append(['Revizija','HYD ID','Premer [mm]','Tlak meja [bar]','Idealna sila pri izbranem tlaku [N]','Tlak dovoljen','COST ID','Sintetična cena [EUR]'])
    for h in results:
        for item in h.results['hidravlik']['engineering']['catalogue']['alternatives']:
            alt.append([h.revision,item['id'],item['bore_mm'],item['max_pressure_bar'],item['ideal_capacity_at_selected_pressure_N'],'da' if item['pressure_permitted'] else 'ne',item['price_id'],item['price_EUR']])
    info=wb.create_sheet('Omejitve');info.append(['Vse cene so izmišljene, izključno iz KB-01. Nobena ponudba ni preverjena.'])
    info.append(['CON-02/03 nimata sledljive geometrije in količin; celotnih konceptnih cen ne določamo.'])
    info.append(['Količina 1 pomeni scenarijski sklop, ne proizvodno kosovnico. COST-10/11/12 ostanejo neznani.'])
    info.append([state['knowledge']['document']+' SHA-256: '+state['knowledge']['sha256']])
    for row in range(1,info.max_row+1):info.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8)
    for sheet in wb:
        for col in 'ABCDEFGH':sheet.column_dimensions[col].width=20 if col not in 'BGH' else 34
        for row in sheet:
            for cell in row:cell.font=Font(name='DejaVu Sans',size=9);cell.alignment=Alignment(wrap_text=True,vertical='top')
        for row in range(1,sheet.max_row+1):sheet.row_dimensions[row].height=60 if sheet==info else 48
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:2'
        sheet.oddFooter.center.text='V01 | Sintetično; tehnično nepotrjeno | Stran &P'
    wb.save(path)


def build_final_package(state,directory):
    # Entire validation/export is offline, including --report-only.
    state=redact(state);results=validate(state);original=results[-1];final,prose_notes=document_view(original);kb=state['knowledge'];directory=Path(directory)
    state['document_presentation']={'model_calls':0,'source_handoff_sha256':digest(original.to_dict()),
        'policy':'Published hydraulic facts derive from validated engineering; original AI prose remains unchanged in revisions.',
        'notes':prose_notes}
    directory.mkdir(parents=True,exist_ok=True);destination=directory/'V01_Koncni_AI_razvojni_paket.zip'
    source_text=json.dumps(state,ensure_ascii=False,indent=2,allow_nan=False)+'\n'
    provenance={'input_file':'08_Izvorni_podatki.json','input_file_sha256':hashlib.sha256(source_text.encode('utf-8')).hexdigest()}
    with TemporaryDirectory(prefix='.final-',dir=directory) as tmp:
        root=Path(tmp);make_pdfs(root,final,provenance,source_filename='08_Izvorni_podatki.json')
        # Add supplements inside the reused PDFs, not detached from their actual stage data.
        from pypdf import PdfReader, PdfWriter
        def supplement(name,title,story):
            path=root/('append-'+name);pdf(path,title,final,story,provenance)
            writer=PdfWriter()
            for source in (root/name,path):
                for page in PdfReader(source).pages:writer.add_page(page)
            with (root/name).open('wb') as stream:writer.write(stream)
            path.unlink()
        summary=[]
        section(summary,'Potek zaključnega preizkusa',[f"Način: {state['mode']}; zaključenih revizij: {len(results)}; API-klicev: {state['api_usage']['call_count']}.",
            f"Ustavitev: {state['stop_reason']}. Baza: {kb['document']}, SHA-256 {kb['sha256']}.",
            'Programski uspeh ni tehnična odobritev. Sintetična konstrukcija ni certificirana ali primerna za izdelavo.',
            'Uporabljene spremembe so navedene v primerjavi; odsotnost spremembe ni prikazana kot izboljšanje.'])
        if prose_notes:section(summary,'Kontrola besedilne skladnosti',['Izvirno besedilo Hidravlika omenja drug katalogski kandidat kot strukturirana izbira. Za tehnične podatke v dokumentih velja preverjena strukturirana predaja; komentarji so nespremenjeni v izvornem JSON.'])
        section(summary,'Odprte in nepotrjene naloge',[t['id']+': '+t['reason'] for t in state['tasks'] if t['status']!='resolved'])
        supplement('00_Povzetek.pdf','Potek razvoja',summary);(root/'00_Povzetek.pdf').rename(root/'00_Povzetek_razvoja.pdf')
        story=[];section(story,'Spremembe geometrije in ponovni izračuni',[])
        for change in state['changes']:
            if change['parameter']!='hydraulic_selection':section(story,change['revision'],[json.dumps(change,ensure_ascii=False)])
        if not any(c['parameter']!='hydraulic_selection' for c in state['changes']):section(story,'Geometrija ohranjena',['Ni utemeljenega sprejetega geometrijskega popravka; deterministični rezultati so ponovno preverjeni.'])
        supplement('01_Konstruktor.pdf','Geometrijske spremembe',story)
        cat=final.results['hidravlik']['engineering']['catalogue'];story=[]
        section(story,'Izbrani sintetični komponenti',[f"Cilinder {cat['choice']['cylinder_id']}, agregat {cat['choice']['pump_id']}, tlak {fmt(cat['choice']['pressure_bar'])} bar.",
            f"Idealna zmogljivost: {fmt(cat['ideal_capacity_N'])} N; zahtevana sila: {fmt(cat['peak_required_force_N'])} N.",cat['basis'],
            'Hod kataloškega cilindra, vgradnja in izkoristek: NI PODATKA. Komponenti sta scenarijska kandidata, ne potrjena izbira.'])
        story.append(table(['HYD ID','D [mm]','p največ [bar]','Idealna sila [N]','COST EUR'],[[a['id'],fmt(a['bore_mm']),fmt(a['max_pressure_bar']),fmt(a['ideal_capacity_at_selected_pressure_N']),fmt(a['price_EUR'])] for a in cat['alternatives']]))
        if prose_notes:section(story,'Kontrola besedilne skladnosti',['Opozorilo: izvirni AI komentar govori tudi o drugem cilindru. Ta komentar ni potrjen tehnični podatek; veljajo zgornja strukturirana izbira in izračuni. Izvirnik in mesta neskladij so v 08_Izvorni_podatki.json.'])
        supplement('02_Hidravlik.pdf','Katalog in primerjava',story)
        story=[];section(story,'Zahtevki in neodvisni statusi',[])
        for t in state['tasks']:section(story,t['id']+' · '+t['status'],[t['reason'],'Odgovoren: '+(t['target_agent'] or 'Ročni pregled'),t.get('resolution','Brez dokaza zaprtja.'),'Viri: '+', '.join(t['source_ids'])])
        if prose_notes:section(story,'Dodatna kontrola izvoznika (ni izvirna ugotovitev AI-recenzenta)',['V hidravličnem besedilu je zaznana omemba drugega cilindra. Neoverjeno besedilo je ohranjeno v JSON in komunikacijskem dnevniku; tehnični pregled dokumentov uporablja strukturirane vrednosti. To ne zapira nobene inženirske ugotovitve.'])
        supplement('04_Recenzent.pdf','Vsi zahtevki',story)
        workbook(root/'03_Ekonomist.xlsx',state,results)
        story=[];section(story,'Primerjava dejansko zaključenih revizij',['Nedokončan predlog ni nova revizija. Manjša idealna sila ali odpravljena kataloška napaka ne dokazujeta splošne tehnične ustreznosti.'])
        story.append(table(['Revizija','F največ [N]','Hod [mm]','Cilinder','Delni COST [EUR]'],[[h.revision,fmt(h.results['konstruktor']['engineering']['peak_cylinder_force_N']),fmt(h.results['konstruktor']['engineering']['stroke_mm']),h.results['hidravlik']['engineering']['catalogue']['choice']['cylinder_id'],fmt(h.results['ekonomist']['engineering']['synthetic_subtotal_EUR'])] for h in results]))
        for change in state['changes']:section(story,'Dokumentirana sprememba',[json.dumps(change,ensure_ascii=False)])
        section(story,'Fiksne zahteve',[json.dumps(state['fixed_requirements'],ensure_ascii=False)])
        section(story,'Odpravljene ugotovitve',[t['id']+': '+t.get('resolution','') for t in state['tasks'] if t['status']=='resolved'] or ['Ni neodvisno dokazanih popravkov.'])
        section(story,'Končna ocena',['changes_required; engineering_verified=false. '+state['stop_reason']])
        pdf(root/'05_Primerjava_revizij.pdf','Primerjava revizij',final,story,provenance)
        story=[];section(story,'Komunikacija prek orkestratorja',['Vsako sporočilo ima pošiljatelja, prejemnika, revizijo, zahtevek, razlog, vir in status. Trditev addressed ne zapre naloge.'])
        if prose_notes:section(story,'Opozorilo za branje izvirne komunikacije',['Spodaj so izvirna sporočila, vključno z nepreverjenim besedilom Hidravlika. Omemba izhodiščnega cilindra ni izbira za končni izračun. Za končno izbiro glej strukturirano predajo in 02_Hidravlik.pdf.'])
        for msg in state['messages']:
            section(story,f"{msg['revision']} · {msg['sender']} → {msg['recipient']} · {msg['request']}",[msg['reason'],'Status: '+msg['status'],'Viri: '+', '.join(msg['data_sources'])])
        pdf(root/'06_Komunikacija_agentov.pdf','Komunikacija agentov',final,story,provenance)
        story=[];section(story,'Vir in omejitve',[kb['document'],'SHA-256: '+kb['sha256'],'Vsa izhodišča, cene in komponente so sintetični podatki KB-01. Polno besedilo in vse tabele so ohranjeni v JSON.'])
        for conflict in kb['conflicts']:section(story,conflict['severity'],[conflict.get('reason',str(conflict))])
        used=set()
        for h in results:
            for name in ORDER:used.update(h.results[name].get('effective_source_ids',h.results[name]['declared_source_ids']))
        for identifier in sorted(used):
            row=kb['records'][identifier];section(story,identifier+' · '+row['classification'],[' | '.join(row['cells']),json.dumps(row['source'],ensure_ascii=False)])
        section(story,'Kaj je bilo dejansko uporabljeno',['Zgornji ID-ji so izrecno navedeni v vrhnjem seznamu ali strukturiranih zahtevkih validiranih odgovorov agentov. Izvorna mesta citatov so shranjena v JSON; novih citatov ne dodajamo. Celoten izvleček je bil posredovan v pozivih; deklaracija ni dokaz notranjega razmišljanja modela.'])
        pdf(root/'07_Viri_in_predpostavke.pdf','Viri in predpostavke',final,story,provenance)
        (root/'08_Izvorni_podatki.json').write_text(source_text,encoding='utf-8')
        (root/'09_API_poraba.json').write_text(json.dumps(state['api_usage'],ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
        with ZipFile(root/destination.name,'w',ZIP_DEFLATED) as archive:
            for name in FILES:archive.write(root/name,name)
        with ZipFile(root/destination.name) as archive:
            if archive.testzip() is not None or set(archive.namelist())!=set(FILES):raise ValueError('Neveljaven končni ZIP.')
        (root/destination.name).replace(destination)
    return destination
